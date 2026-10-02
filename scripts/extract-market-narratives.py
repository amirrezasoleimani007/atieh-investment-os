"""Reproducible import of approved company prose and every Trade row.

Usage: python scripts/extract-market-narratives.py ../upload [--check]
Only extracts sources; does not infer financial or market figures.
"""
import hashlib
import json
import re
import sys
from pathlib import Path
from docx import Document
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
inputs = Path(sys.argv[1])
check = '--check' in sys.argv

def save(name, data):
    target = root / 'public/data' / name
    if check:
        assert json.loads(target.read_text()) == data, name
    else:
        target.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n')

def provenance(path):
    return {'file': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}

doc = inputs / 'final.docx'
management = []
for paragraph in Document(doc).paragraphs:
    text = paragraph.text
    if not text.strip():
        continue
    if paragraph.style.name == 'Heading 2' and re.match(r'^[\d۰-۹]+[.٫]', text):
        management.append({'name': re.sub(r'^[\d۰-۹]+[.٫]\s*', '', text).strip(), 'title': '', 'paragraphs': []})
    elif management:
        if paragraph.style.name == 'Heading 3':
            management[-1]['title'] = text
        else:
            management[-1]['paragraphs'].append(text)
assert len(management) == 10
save('management-narratives.json', {'source': provenance(doc), 'companies': management})

md = inputs / 'Pasted markdown(20260923-083551).md'
companies = []
for block in re.split(r'^# ', md.read_text(), flags=re.M)[1:]:
    lines = block.splitlines()
    name = re.sub(r'^[\d۰-۹]+\.\s*', '', lines[0]).strip()
    sections = []
    for line in lines[1:]:
        if not line.strip() or line == '---':
            continue
        if line.startswith('## '):
            sections.append({'heading': line[3:], 'blocks': []})
        else:
            if not sections:
                sections.append({'heading': '', 'blocks': []})
            sections[-1]['blocks'].append({'type': 'heading' if line.startswith('### ') else 'paragraph', 'text': line[4:] if line.startswith('### ') else line})
    companies.append({'id': f'company-{len(companies)+1}', 'name': name, 'sections': sections})
assert len(companies) == 10
# Explicit links grounded in the supplied company narratives; not fuzzy name matching.
mapping = {'steel': [1, 2, 3, 4, 7, 9], 'welded': [8], 'seamless': [10], 'cored': [6], 'bentonite': [6], 'paint': [5]}
save('company-market-narratives.json', {'source': provenance(md), 'companies': companies, 'marketCompanies': {key: [f'company-{i}' for i in ids] for key, ids in mapping.items()}})

excel = inputs / 'داشبورد(1).xlsx'
workbook = load_workbook(excel, data_only=True)
market_rows = []
for rowno, row in enumerate(workbook['Market'].iter_rows(min_row=2, values_only=True), 2):
    if row[0] and row[3]:
        market_rows.append({'row': rowno, 'scope': row[0], 'product': row[3], 'scenario': row[5], 'companies': [name for name in row[10:14] if name]})
save('market-company-links.json', {'source': provenance(excel), 'records': market_rows})
series = {}
records = []
for rowno, row in enumerate(workbook['Trade'].iter_rows(min_row=2, values_only=True), 2):
    year, group, product, value, *related = row
    if product is None:
        continue
    assert isinstance(value, (int, float)), (rowno, value)
    product = product.strip()
    item = series.setdefault(product, {'product': product, 'group': group, 'values': [], 'companies': []})
    item['values'].append(value)
    for name in related:
        if name and name not in item['companies']:
            item['companies'].append(name)
    records.append({'row': rowno, 'year': year, 'product': product, 'value': value, 'companies': [x for x in related if x]})
years = list(dict.fromkeys(str(x['year']) for x in records))
assert len(series) == 8 and len(records) == 56
assert all(len(x['values']) == len(years) for x in series.values())
save('steel-trade.json', {'source': provenance(excel), 'unit': 'هزار تن', 'years': years, 'series': list(series.values()), 'records': records})
print(f"{'Verified' if check else 'Imported'}: 10 management companies; 10 market narratives; 8 products / 56 trade records")
