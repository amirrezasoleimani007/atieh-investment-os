from __future__ import annotations

import json
import math
import re
import zipfile
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


UPLOAD = Path("/workspace/scratch/7f1d59fd8c35/upload")
OUT = Path(__file__).resolve().parents[1] / "public" / "data"
FINAL_FIT_WORKBOOK = UPLOAD / "مدل_نهایی_Fit_علمی_77_بخش_نسخه_2_2(1)(1).xlsx"
FINAL_IO_WORKBOOK = UPLOAD / "ارتباط_پرتفوی_IO_قابلیت_و_زنجیره_77_بخش_نسخه_نهایی_2_2(2).xlsx"


def clean(value: object) -> str:
    text = "" if value is None else str(value)
    return re.sub(r"\s+", " ", text.replace("ي", "ی").replace("ك", "ک").replace("\u200c", " ")).strip()


def number(value: object) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return float(value)
    return None


def write_json(name: str, payload: object) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / name).write_text(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )


def records_by_header(ws, header_row: int, first_data_row: int, last_data_row: int, max_col: int | None = None) -> dict[int, dict[str, object]]:
    """Read governed rows by normalized header, never by legacy column position."""
    headers = [clean(ws.cell(header_row, column).value) for column in range(1, (max_col or ws.max_column) + 1)]
    if len(headers) != len(set(header for header in headers if header)):
        raise ValueError(f"Duplicate headers in {ws.title}")
    records: dict[int, dict[str, object]] = {}
    for row in ws.iter_rows(min_row=first_data_row, max_row=last_data_row, max_col=max_col, values_only=True):
        record = {header: value for header, value in zip(headers, row, strict=False) if header}
        row_id = number(record.get("ردیف"))
        if row_id is not None:
            records[int(row_id)] = record
    return records


def extract_sectors() -> list[dict]:
    topsis = load_workbook(
        UPLOAD / "TOPSIS_Dynamic_Opportunity_Ranking(2).xlsx",
        read_only=True,
        data_only=True,
    )["Sheet1"]
    architecture = load_workbook(
        UPLOAD / "IPS_77_New_Architecture_Final(2)(2).xlsx",
        read_only=True,
        data_only=True,
    )["ورودی"]
    top_rows = list(topsis.iter_rows(min_row=2, max_row=78, max_col=7, values_only=True))
    arch_rows = list(architecture.iter_rows(min_row=2, max_row=78, max_col=7, values_only=True))
    sectors: list[dict] = []
    for top, arch in zip(top_rows, arch_rows, strict=True):
        if int(top[0]) != int(arch[0]):
            raise ValueError("Sector row alignment mismatch")
        sectors.append(
            {
                "id": int(top[0]),
                "name": clean(top[1]),
                "criteria": {
                    "growth": number(top[2]),
                    "valueAdded": number(top[3]),
                    "megatrends": number(top[4]),
                    "inflation": number(top[5]),
                    "fxExposure": number(top[6]),
                },
                "baselineOpportunity": number(arch[2]),
                "relatedness": {
                    "financial": number(arch[3]),
                    "downstream": number(arch[4]),
                    "trade": number(arch[5]),
                    "support": number(arch[6]),
                },
            }
        )
    return sectors


DIMENSIONS = {
    "growth": {"label": "رشد", "weight": 0.20, "codes": {"144": 0.40, "186": 0.45, "187": 0.15}},
    "profitability": {"label": "سودآوری", "weight": 0.15, "codes": {"146": 1.0}},
    "capitalReturn": {"label": "ارزش‌آفرینی", "weight": 0.20, "codes": {"176": 1.0}},
    "cashQuality": {"label": "کیفیت سود", "weight": 0.15, "codes": {"162": 0.40, "175": 0.40, "169": 0.20}},
    "resilience": {"label": "تاب‌آوری", "weight": 0.20, "codes": {"174": 0.50, "173": 0.30, "164": 0.20}},
    "workingCapital": {"label": "بهره‌وری", "weight": 0.10, "codes": {"152": 0.60, "168": 0.40}},
}


def dimension_score(records: dict[str, dict], definition: dict) -> tuple[float | None, float]:
    coverage = sum(
        weight
        for code, weight in definition["codes"].items()
        if records.get(code, {}).get("score") is not None
    )
    if coverage + 1e-9 < 0.60:
        return None, coverage
    score = sum(
        records[code]["score"] * weight / coverage
        for code, weight in definition["codes"].items()
        if records.get(code, {}).get("score") is not None
    )
    return round(score, 4), round(coverage, 4)


def annual_score(records: dict[str, dict]) -> dict:
    dimensions = {}
    for key, definition in DIMENSIONS.items():
        score, coverage = dimension_score(records, definition)
        dimensions[key] = {"score": score, "coverage": coverage, "label": definition["label"]}
    available_weight = sum(
        definition["weight"]
        for key, definition in DIMENSIONS.items()
        if dimensions[key]["score"] is not None
    )
    score = None
    if available_weight + 1e-9 >= 0.60:
        score = sum(
            dimensions[key]["score"] * definition["weight"] / available_weight
            for key, definition in DIMENSIONS.items()
            if dimensions[key]["score"] is not None
        )
        score = round(score, 4)
    return {
        "score": score,
        "coverage": round(available_weight, 4),
        "dimensions": dimensions,
        "validity": "valid" if available_weight >= 0.75 else "temporary" if available_weight >= 0.60 else "insufficient",
    }


def performance_label(score: float | None) -> str:
    if score is None:
        return "داده ناکافی"
    if score >= 8:
        return "ممتاز"
    if score >= 6.5:
        return "مطلوب"
    if score >= 5:
        return "متوسط"
    if score >= 3.5:
        return "ضعیف"
    return "بحرانی"


def extract_swot() -> dict[str, dict]:
    ns = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    }
    company_by_slide = {
        1: "آتیه تجارت نقش جهان قشم",
        2: "آتیه صنعت افق نقش جهان",
        3: "نورد لوله کوثر صنعت اسپادانا",
        4: "فولاد متیل",
        5: "ورق خودرو چهارمحال بختیاری",
        6: "کارگزاری مبین سرمایه",
        7: "صنایع برش ورق فولادی مبارکه",
        8: "پولای بهیز",
        9: "توکا رنگ فولاد سپاهان",
    }
    result: dict[str, dict] = {}
    path = UPLOAD / "14050417 - SWOT(2).pptx"
    with zipfile.ZipFile(path) as archive:
        slide_names = sorted(
            (name for name in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)),
            key=lambda name: int(re.search(r"slide(\d+)\.xml$", name).group(1)),
        )
        for name in slide_names:
            index = int(re.search(r"slide(\d+)\.xml$", name).group(1))
            root = ET.fromstring(archive.read(name))
            blocks: list[list[str]] = []
            for shape in root.findall(".//p:sp", ns):
                lines = []
                for paragraph in shape.findall(".//a:p", ns):
                    text = clean("".join(node.text or "" for node in paragraph.findall(".//a:t", ns))).lstrip("•").strip()
                    if text:
                        lines.append(text)
                if lines:
                    blocks.append(lines)
            if len(blocks) < 9:
                raise ValueError(f"Unexpected SWOT structure on slide {index}")
            result[company_by_slide[index]] = {
                "strengths": blocks[0],
                "weaknesses": blocks[2],
                "threats": blocks[5],
                "opportunities": blocks[8],
                "sourceNote": "نام شرکت در اسلایدهای ۸ و ۹ فایل مرجع با محتوای فعالیت تطبیق و اصلاح شده است."
                if index in (8, 9)
                else "برگرفته از فایل SWOT مصوب",
            }
    return result


def extract_market() -> dict[str, dict]:
    ws = load_workbook(UPLOAD / "3.market .xlsx", read_only=True, data_only=True).active
    result: dict[str, dict] = defaultdict(dict)
    for company, section, score in ws.iter_rows(min_row=2, max_col=3, values_only=True):
        if company and section and number(score) is not None:
            result[clean(company)][clean(section)] = round(float(score), 4)
    return dict(result)


def extract_companies() -> list[dict]:
    ws = load_workbook(UPLOAD / "1.Atieh(2).xlsx", read_only=True, data_only=True)["دیتا بیس"]
    by_company: dict[str, dict[int, dict[str, dict]]] = defaultdict(lambda: defaultdict(dict))
    nature_by_company: dict[str, str] = {}
    variables: dict[str, str] = {}
    bad_values: dict[str, int] = defaultdict(int)
    total_rows: dict[str, int] = defaultdict(int)
    score_rows: dict[str, int] = defaultdict(int)
    valid_score_rows: dict[str, int] = defaultdict(int)
    for row in ws.iter_rows(min_row=2, max_col=9, values_only=True):
        nature, company, year, variable, code, unit, value, score, _mark = row
        if not company or not year or code is None:
            continue
        company_name = clean(company)
        code_text = str(int(code)) if isinstance(code, float) and code.is_integer() else str(code)
        total_rows[company_name] += 1
        nature_by_company[company_name] = clean(nature)
        variables[code_text] = clean(variable)
        parsed_value = number(value)
        parsed_score = number(score)
        if score is not None:
            score_rows[company_name] += 1
        if parsed_score is not None and 0 <= parsed_score <= 10:
            valid_score_rows[company_name] += 1
        else:
            parsed_score = None
        if value is not None and parsed_value is None:
            bad_values[company_name] += 1
        by_company[company_name][int(year)][code_text] = {
            "value": parsed_value,
            "score": parsed_score,
            "unit": clean(unit),
            "variable": clean(variable),
        }

    market = extract_market()
    swot = extract_swot()
    kpi_codes = ["111", "118", "124", "114", "134", "136", "146", "153", "176", "168", "164", "173"]
    companies = []
    for company, years_map in by_company.items():
        years = sorted(years_map)
        annual = {str(year): annual_score(years_map[year]) for year in years}
        current_year = max(years)
        current = annual[str(current_year)]
        prior_scores = [annual.get(str(current_year - offset), {}).get("score") for offset in (0, 1, 2)]
        if prior_scores[0] is None:
            history_score = None
            history_label = "غیرقابل محاسبه"
        elif prior_scores[1] is not None and prior_scores[2] is not None:
            history_score = round(0.5 * prior_scores[0] + 0.3 * prior_scores[1] + 0.2 * prior_scores[2], 4)
            history_label = "سه‌ساله ۵۰٪، ۳۰٪ و ۲۰٪"
        elif prior_scores[1] is not None:
            history_score = round((0.5 * prior_scores[0] + 0.3 * prior_scores[1]) / 0.8, 4)
            history_label = "دوساله با بازتوزیع"
        else:
            history_score = prior_scores[0]
            history_label = "مبتنی بر یک سال"
        kpis = {}
        for code in kpi_codes:
            current_record = years_map[current_year].get(code)
            previous_record = years_map.get(current_year - 1, {}).get(code)
            if current_record:
                kpis[code] = {
                    "label": variables.get(code, code),
                    "value": current_record["value"],
                    "previous": previous_record["value"] if previous_record else None,
                    "score": current_record["score"],
                    "unit": current_record["unit"],
                }
        market_data = market.get(company, {})
        if not market_data:
            company_key = clean(company).replace(" و ", " ")
            market_data = next(
                (value for name, value in market.items() if clean(name).replace(" و ", " ") == company_key),
                {},
            )
        companies.append(
            {
                "name": company,
                "nature": nature_by_company[company],
                "year": current_year,
                "score": current["score"],
                "historyScore": history_score,
                "historyLabel": history_label,
                "classification": performance_label(current["score"]),
                "coverage": current["coverage"],
                "validity": current["validity"],
                "dimensions": current["dimensions"],
                "annualScores": {year: annual[year]["score"] for year in annual},
                "kpis": kpis,
                "market": market_data,
                "swot": swot.get(company, {"strengths": [], "weaknesses": [], "opportunities": [], "threats": [], "sourceNote": "داده SWOT موجود نیست"}),
                "dataQuality": {
                    "records": total_rows[company],
                    "invalidValues": bad_values[company],
                    "scoreCompleteness": round(valid_score_rows[company] / score_rows[company], 4) if score_rows[company] else 0,
                },
            }
        )
    companies.sort(key=lambda item: item["score"] if item["score"] is not None else -1, reverse=True)
    last_score = None
    rank = 0
    for index, item in enumerate(companies, start=1):
        if item["score"] is None:
            item["rank"] = None
            continue
        if last_score is None or abs(item["score"] - last_score) > 1e-9:
            rank = index
        item["rank"] = rank
        last_score = item["score"]
    return companies


def extract_isic() -> list[dict]:
    ws = load_workbook(UPLOAD / "ISIC Mapping.xlsx", read_only=True, data_only=True)["DATABASE"]
    rows = []
    for description, code, sector_code, sector in ws.iter_rows(min_row=2, max_col=4, values_only=True):
        if not description:
            continue
        code_text = ""
        if code is not None:
            code_text = str(int(code)) if isinstance(code, float) and code.is_integer() else clean(code)
        sector_code_text = ""
        if sector_code is not None:
            sector_code_text = str(int(sector_code)) if isinstance(sector_code, float) and sector_code.is_integer() else clean(sector_code)
        rows.append(
            {
                "description": clean(description),
                "isic": code_text,
                "sectorCode": sector_code_text,
                "sector": clean(sector),
            }
        )
    return rows


def extract_portfolio() -> list[dict]:
    ws = load_workbook(UPLOAD / "درصد سهامداری کل.xlsx", read_only=True, data_only=True).active
    rows = []
    for company, assets, ownership, attributable, share, nature, status in ws.iter_rows(
        min_row=2, max_row=12, max_col=7, values_only=True
    ):
        if not company:
            continue
        nature_clean = clean(nature)
        rows.append(
            {
                "name": clean(company),
                "assets": number(assets),
                "ownership": number(ownership),
                "attributableValue": number(attributable),
                "portfolioShare": number(share),
                "nature": nature_clean,
                "control": "کنترلی" if clean(status).replace(" ", "") in {"کنترلی", "کنترلیی"} else "غیرکنترلی",
                "horizon": "هسته" if "بازرگانی" in nature_clean else "مجاور",
            }
        )
    return rows


def extract_growth_opportunities() -> tuple[list[dict], dict[str, object]]:
    wb = load_workbook(FINAL_FIT_WORKBOOK, read_only=True, data_only=True)
    io_wb = load_workbook(FINAL_IO_WORKBOOK, read_only=True, data_only=True)
    fit_rows = records_by_header(wb["محاسبات Fit"], 3, 6, 82, 21)
    extra_rows = records_by_header(wb["تحلیل تکمیلی"], 4, 5, 81, 16)
    io_rows = records_by_header(io_wb["خروجی IO"], 3, 4, 80, 16)
    io_base_rows = records_by_header(wb["IO ارتباط پایه"], 3, 4, 80, 34)
    network_ws = io_wb["شبکه Transform 77"]
    network_rows = records_by_header(network_ws, 3, 4, 80, 81)
    expected_ids = set(range(1, 78))
    if any(set(rows) != expected_ids for rows in (fit_rows, extra_rows, io_rows, io_base_rows, network_rows)):
        raise ValueError("Fit, detailed-analysis, IO and network rows are not aligned to 77 opportunities")

    adjacent_headers = {
        "downstream": "Eff پایین دستی",
        "financeDirect": "Eff مالی مستقیم",
        "mineral": "Eff معدنی",
        "steelSupport": "Eff جانبی",
        "paint": "Eff رنگ",
    }
    result = []
    for idx in sorted(fit_rows):
        row = fit_rows[idx]
        extra = extra_rows[idx]
        io_row = io_rows[idx]
        io_base = io_base_rows[idx]
        names = (row["بخش"], extra["فرصت"], io_row["بخش"], io_base["بخش"])
        if len({clean(value) for value in names}) != 1:
            raise ValueError(f"Opportunity name mismatch at row {idx}")
        adjacent_effective = {key: number(row[header]) for key, header in adjacent_headers.items()}
        core_fit = number(row["Fit هسته بازرگانی فولاد"])
        adjacent_fit = number(row["Adjacent Fit"])
        transform_fit = number(row["Transform Fit – Network Density"])
        expected = (core_fit, *adjacent_effective.values(), adjacent_fit)
        actual = (
            number(io_row["Fit هسته"]), number(io_row["پایین دستی"]), number(io_row["مالی مستقیم"]),
            number(io_row["معدنی"]), number(io_row["صنایع جانبی"]), number(io_row["رنگ"]), number(io_row["Adjacent Fit"]),
        )
        if any(a is None or b is None or not math.isclose(a, b, rel_tol=1e-12, abs_tol=1e-12) for a, b in zip(expected, actual, strict=True)):
            raise ValueError(f"Fit output mismatch between reference workbooks at row {idx}")
        result.append(
            {
                "id": idx,
                "name": clean(row["بخش"]),
                "ioCode": int(io_row["کد IO"]),
                "ioActivity": clean(io_base["فعالیت IO"]),
                "opportunityRaw": number(row["TOPSIS خام"]),
                "xPlotBaseline": number(row["X نمودار – مقیاس کشیده"]),
                "coreFit": core_fit,
                "adjacentEffective": adjacent_effective,
                "adjacentFit": adjacent_fit,
                "adjacentDriver": clean(io_row["مسیر مجاور غالب"]),
                "financeEnablement": number(io_row["توانمندسازی مالی"]),
                "transformFit": transform_fit,
                "yRawBaseline": number(row["Y خام مبنا"]),
                "yPlotBaseline": number(row["Y نمودار مبنا – مقیاس کشیده"]),
                "closest": clean(extra["نزدیک ترین بخش پرتفوی"]),
                "closeness": number(extra["Closeness 1"]),
                "secondClosest": clean(extra["دومین بخش نزدیک"]),
                "secondCloseness": number(extra["Closeness 2"]),
                "gap": number(extra["Gap"]),
                "synergyPair": clean(extra["Synergy Pair"]),
                "synergy": number(extra["Synergy Potential Index"]),
                "optionalityRaw": number(extra["Optionality Raw"]),
                "optionality": number(extra["Optionality Index"]),
                "nextPaths": [clean(extra[f"مسیر بعدی {path}"]) for path in (1, 2, 3)],
                "verticalTrade": number(io_row["Vertical تجارت"]),
                "verticalSteel": number(io_row["Vertical فولاد"]),
                "verticalAdjacent": number(io_row["Vertical مجاور"]),
                "leontiefTotal": number(io_row["Leontief کل زنجیره"]),
            }
        )

    network_headers = [clean(network_ws.cell(3, column).value) for column in range(5, 82)]
    names = [row["name"] for row in result]
    if network_headers != names:
        raise ValueError("Structural-relatedness destination order does not match opportunity order")
    matrix = []
    for idx in range(1, 78):
        values = [number(network_rows[idx][name]) for name in network_headers]
        if any(value is None for value in values):
            raise ValueError(f"Missing structural-relatedness value at row {idx}")
        matrix.append(values)
    relatedness = {"model": "2.2", "ids": list(range(1, 78)), "names": names, "matrix": matrix}
    return result, relatedness


def extract_financial_scope(filename: str, scope: str) -> list[dict]:
    ws = load_workbook(UPLOAD / filename, read_only=True, data_only=True)["دیتا بیس"]
    by_company: dict[str, dict[int, dict[str, dict]]] = defaultdict(lambda: defaultdict(dict))
    nature_by_company: dict[str, str] = {}
    variables: dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, max_col=9, values_only=True):
        nature, company, year, variable, code, unit, value, score, _mark = row
        if not company or not year or code is None:
            continue
        company_name = clean(company)
        code_text = str(int(code)) if isinstance(code, float) and code.is_integer() else str(code)
        nature_by_company[company_name] = clean(nature)
        variables[code_text] = clean(variable)
        parsed_score = number(score)
        if parsed_score is None or not 0 <= parsed_score <= 10:
            parsed_score = None
        by_company[company_name][int(year)][code_text] = {
            "value": number(value),
            "score": parsed_score,
            "unit": clean(unit),
        }
    # Preserve the complete financial-variable catalog from the governed source.
    # The interface can expose every available actual indicator instead of a
    # hand-picked subset.
    kpi_codes = sorted(variables, key=lambda code: int(code) if code.isdigit() else 10_000)
    result = []
    for company, years_map in by_company.items():
        years = sorted(years_map)
        current_year = max(years)
        current_analysis = annual_score(years_map[current_year])
        yearly_analysis = {str(year): annual_score(years_map[year]) for year in years}
        annual = {year: analysis["score"] for year, analysis in yearly_analysis.items()}
        scatter_specs = {
            "scatterOperationalGrowth": [("144", 0.55), ("186", 0.45)],
            "scatterOperatingProfitability": [("146", 0.60), ("178", 0.40)],
            "scatterResourceProductivity": [("152", 0.60), ("154", 0.40)],
            "scatterCapitalValueCreation": [("176", 0.60), ("151", 0.40)],
            "scatterLiquidityStrength": [("164", 0.60), ("163", 0.40)],
            "scatterFinancialStructure": [("173", 0.40), ("174", 0.35), ("155", 0.25)],
            "scatterEarningsQuality": [("162", 0.60), ("169", 0.40)],
            "scatterFcfGeneration": [("175", 0.60), ("170", 0.40)],
        }
        scatter_axes = {}
        for axis, components in scatter_specs.items():
            values = [(years_map[current_year].get(code) or {}).get("score") for code, _ in components]
            scatter_axes[axis] = (
                round(sum(value * weight for value, (_, weight) in zip(values, components)), 4)
                if all(value is not None for value in values)
                else None
            )
        kpis = {}
        kpi_history = {}
        for code in kpi_codes:
            current = years_map[current_year].get(code)
            previous = years_map.get(current_year - 1, {}).get(code)
            if current:
                kpis[code] = {
                    "label": variables.get(code, code),
                    "value": current["value"],
                    "previous": previous["value"] if previous else None,
                    "score": current["score"],
                    "unit": current["unit"],
                }
            history = []
            for year in years:
                record = years_map[year].get(code)
                if record and record["value"] is not None:
                    history.append({
                        "year": year,
                        "value": record["value"],
                        "score": record["score"],
                        "unit": record["unit"],
                        "label": variables.get(code, code),
                    })
            if history:
                kpi_history[code] = history
        result.append(
            {
                "name": company,
                "scope": scope,
                "nature": nature_by_company[company],
                "year": current_year,
                "score": current_analysis["score"],
                "classification": performance_label(current_analysis["score"]),
                "coverage": current_analysis["coverage"],
                "dimensions": current_analysis["dimensions"],
                "annualScores": annual,
                "yearlyDimensions": {year: analysis["dimensions"] for year, analysis in yearly_analysis.items()},
                "scatterAxes": scatter_axes,
                "kpis": kpis,
                "kpiHistory": kpi_history,
            }
        )
    result.sort(key=lambda item: item["score"] if item["score"] is not None else -1, reverse=True)
    for rank, item in enumerate([item for item in result if item["score"] is not None], start=1):
        item["rank"] = rank
    for item in result:
        item.setdefault("rank", None)
    return result


def main() -> None:
    sectors = extract_sectors()
    companies = extract_companies()
    isic = extract_isic()
    portfolio = extract_portfolio()
    opportunities, opportunity_relatedness = extract_growth_opportunities()
    financial_companies = extract_financial_scope("1.Atieh(2).xlsx", "گروه آتیه") + extract_financial_scope("2.Metil.xlsx", "گروه متیل")
    write_json("sectors.json", sectors)
    write_json("companies.json", companies)
    write_json("isic.json", isic)
    write_json("portfolio.json", portfolio)
    write_json("opportunities.json", opportunities)
    write_json("opportunity-relatedness.json", opportunity_relatedness)
    write_json("financial-companies.json", financial_companies)
    write_json(
        "manifest.json",
        {
            "snapshot": "۱۴۰۵/۰۴/۱۷",
            "sectors": len(sectors),
            "companies": len(companies),
            "financialCompanies": len(financial_companies),
            "portfolioCompanies": len(portfolio),
            "isicRows": len(isic),
            "sources": [
                "IPS 77 New Architecture",
                "TOPSIS Dynamic Opportunity Ranking",
                "Atieh Financial Database",
                "Market Assessment",
                "ISIC Mapping",
                "SWOT 1405/04/17",
                "مدل نهایی Fit علمی 77 بخش – نسخه 2.2",
                "ارتباط پرتفوی IO، قابلیت و زنجیره 77 بخش – نسخه نهایی 2.2",
                "روش‌شناسی دقیق مدل فرصت و هم‌راستایی استراتژیک – نسخه نهایی 2.2",
            ],
            "model": "IPS-FM-3.0",
            "strategicFitModel": "2.2",
            "strategicFitIoBaseYear": 1400,
        },
    )
    print(f"Generated {len(sectors)} sectors, {len(companies)} companies, {len(isic)} ISIC rows")


if __name__ == "__main__":
    main()
